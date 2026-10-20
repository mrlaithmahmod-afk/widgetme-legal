"""Write financial_model_summary.md from the RECALCULATED workbook so numbers match exactly."""
from openpyxl import load_workbook

XL = "/home/user/widgetme-legal/reports/Property_Visualisation_Financial_Model.xlsx"
OUT = "/home/user/widgetme-legal/research_notes/AI property visualisation SaaS audit/financial_model_summary.md"
wb = load_workbook(XL, data_only=True)


def find(ws, label, start=1, col=1):
    ws = wb[ws]
    for r in range(start, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.strip().startswith(label.strip()):
            return r
    raise KeyError(label)


def v(ws, r, c):
    return wb[ws].cell(r, c).value


def gbp(x, d=0):
    if isinstance(x, str):
        return x
    s = f"£{abs(x):,.{d}f}"
    return f"−{s}" if x < 0 else s


def usd(x, d=2):
    return f"${x:,.{d}f}"


def pct(x):
    return f"{x*100:.1f}%" if not isinstance(x, str) else x


def num(x, d=1):
    return f"{x:,.{d}f}" if not isinstance(x, str) else x


def mon(x):
    return f"M{int(x)}" if isinstance(x, (int, float)) else x


def arow(label):
    r = find("Assumptions", label)
    return [v("Assumptions", r, c) for c in (3, 4, 5)]


L = []
w = L.append
w("# Financial model summary — AI property visualisation SaaS (UK-first)")
w("")
w("Workbook: `/home/user/widgetme-legal/reports/Property_Visualisation_Financial_Model.xlsx`. It has 36 monthly periods and every output is a live formula driven by the `Assumptions` sheet. "
  "Values below were read back after a LibreOffice headless recalculation, which found 0 formula errors. Built 20 Oct 2026. Currency is GBP at **£0.755 per US$** (X-Rates Oct 2026 average). "
  "All inputs come from the research notes in this folder; where a figure was missing the model labels it ASSUMPTION.")
w("")
w("## Bottom line")
br = {k: find("Summary", k) for k in ["Subscription MRR", "ARR (£)", "Paying customers", "Gross margin (month)", "  memo: GM", "Monthly EBITDA", "Cash balance", "Headcount"]}
sl = {k: find("Summary", k) for k in ["First month with EBITDA", "Sustained break-even", "Lowest cash", "Month of lowest", "External funding", "First month cash goes", "Cumulative revenue", "Gross margin, 36", "First month human", "Customers in that", "Month first support", "Month Engineer", "Month AE", "Month first Enterprise"]}
S = lambda key, col: v("Summary", sl[key], col)
B = lambda key, col: v("Summary", br[key], col)
w(f"- **Base case:** month-36 subscription MRR is {gbp(B('Subscription MRR',7))} (ARR {gbp(B('ARR (£)',7))}) from {num(B('Paying customers',7))} customers. "
  f"EBITDA first turns positive in month {S('First month with EBITDA',3)} and stays positive from month {S('Sustained break-even',3)}. "
  f"The £50k founder cash runs out in month {S('First month cash goes',3)}, and the cash low point is {gbp(S('Lowest cash',3))} in month {S('Month of lowest',3)}. "
  f"That means about **{gbp(S('External funding',3))} of external funding** on top of the £50k.")
w(f"- **Bear case:** month-36 ARR is {gbp(B('ARR (£)',4))} and the business never breaks even in 36 months. Gross margin is about {pct(B('Gross margin (month)',4))} under pessimistic unit costs, and the funding need is {gbp(S('External funding',2))} and still growing.")
w(f"- **Bull case:** month-36 ARR is {gbp(B('ARR (£)',10))}. Sustained break-even comes in month {S('Sustained break-even',4)}, the funding need is {gbp(S('External funding',4))}, and month-36 cash is {gbp(B('Cash balance',10))}.")
w(f"- **Verdict:** the idea is viable but not easy. On base assumptions it is a ~£1m-ARR UK business by year 3 that needs about {gbp(50000+S('External funding',3))} in total funding (£50k founder plus {gbp(S('External funding',3))} external). The base gross margin of ~70% holds only if human 2D-confirm/QA stays at ~10 minutes per plan at offshore rates. "
  "Under the pessimistic cost case (15-minute confirms, $35/h QA, more retakes, 50% pixel-streamed tours) every tier falls to roughly a 1–33% gross margin at full included usage, and the Bear scenario never recovers. "
  "So the business depends on QA-minutes-per-plan and the cost of the interactive viewer, not on LLM or image-model prices. Measure both in pilots before committing.")
w("")
w("## Key assumptions (Base unless shown as Bear / Base / Bull)")
fx = arow("FX rate")[1]
rows = [
    ("Starting cash", gbp(arow("Starting cash")[1]), "ASSUMPTION (founder-funded)"),
    ("Image gen per attempt (Low/Base/High)", " / ".join(usd(x, 3) for x in arow("Image generation per attempt")), "VERIFIED Gemini Nano Banana 2.1 / Pro 4K"),
    ("Attempts per accepted image", " / ".join(str(x) for x in arow("Generation attempts")), "ESTIMATE"),
    ("Standard / cinematic video $/s (Base)", f"{usd(arow('Standard walkthrough video')[1])} / {usd(arow('Cinematic video')[1])}", "VERIFIED Veo 3.1 Fast / Standard"),
    ("Human QA rate $/h (Low/Base/High)", " / ".join(usd(x, 0) for x in arow("Human QA / 2D-confirm")), "ESTIMATE (Upwork/offshore)"),
    ("Human 2D confirm, min per plan", " / ".join(str(x) for x in arow("Human 2D vector")), "ESTIMATE (technical_feasibility §7)"),
    ("Deliverable rejection rate", " / ".join(pct(x) for x in arow("Deliverable-level QA")), "ASSUMPTION"),
    ("Pixel-streaming h per hosted property-month", " / ".join(str(x) for x in arow("Interactive viewer")), "ESTIMATE; WebGL default"),
    ("WorkOS SSO", usd(arow("WorkOS SSO price")[1], 0) + "/connection/month", "VERIFIED"),
    ("ODA Drawings SDK (Sustaining)", f"{usd(arow('ODA Drawings SDK')[1],0)} then {usd(arow('ODA Sustaining renewal')[1],0)}/yr", "VERIFIED (reconfirm)"),
    ("ISO 27001", f"{gbp(arow('ISO 27001 year-1')[1])} yr 1, {gbp(arow('ISO 27001 ongoing')[1])}/yr after", "ESTIMATE"),
    ("Insurance (cyber+PI+PL)", gbp(arow("Insurance")[1]) + "/yr", "ESTIMATE"),
    ("Monthly churn Studio / Agency / Dev Pro / Enterprise (Base)",
     " / ".join(pct(arow(f"{t}: monthly logo churn")[1]) for t in ["Studio", "Agency", "Developer Pro", "Enterprise"]), "ESTIMATE (SaaS Capital GRR ~91%/95%; SMB 2.5–3.5%)"),
    ("Launch month Studio / Agency / Dev Pro / Enterprise (Bear / Base / Bull)",
     "; ".join("/".join(str(int(x)) for x in arow(f"{t}: first month")) for t in ["Studio", "Agency", "Developer Pro", "Enterprise"]), "ASSUMPTION (no enterprise before M12 in Base)"),
    ("Included-credit utilisation (Bear/Base/Bull)", " / ".join(pct(x) for x in arow("Average utilisation")), "ASSUMPTION"),
    ("Founder salary; employer on-costs", f"{gbp(arow('Founder salary')[1])}/yr; {pct(arow('Employer on-costs')[1])}", "ASSUMPTION; ESTIMATE"),
    ("Bear scenario cost case", "High (pessimistic); Base and Bull use Base costs", "ASSUMPTION (conservative)"),
]
w("| Input | Value | Label |")
w("|---|---|---|")
for a_, b_, c_ in rows:
    w(f"| {a_} | {b_} | {c_} |")
w("")
hr0 = find("Assumptions", "Role")
w("Hiring is a set of step functions. Each role starts the month after its trigger is first met, measured on the previous month's value: "
  + "; ".join(f"{v('Assumptions', r, 1)} at {'£' if v('Assumptions', r, 2)=='MRR' else ''}{v('Assumptions', r, 3):,} {'MRR' if v('Assumptions', r, 2)=='MRR' else 'customers'}"
              for r in range(hr0 + 1, hr0 + 11)) + ".")
w("")
w("## Unit cost per successful deliverable")
w("Each deliverable cost is built from these components: LLM tokens, image gen, video gen, GPU, streaming, storage, CDN, CAD fees, human 2D confirm, human output QA and support allocation. It includes the case's rejection rate.")
w("")
ur = find("Unit_Costs", "Cost per successful deliverable — summary") + 2
w("| Deliverable | Low £ | Base £ | High £ | Low $ | Base $ | High $ | Base human-labour share |")
w("|---|---|---|---|---|---|---|---|")
for r in range(ur, ur + 11):
    lab = v("Unit_Costs", r, 1)
    w(f"| {lab} | {gbp(v('Unit_Costs',r,5),2)} | {gbp(v('Unit_Costs',r,6),2)} | {gbp(v('Unit_Costs',r,7),2)} | "
      f"{usd(v('Unit_Costs',r,2))} | {usd(v('Unit_Costs',r,3))} | {usd(v('Unit_Costs',r,4))} | {pct(v('Unit_Costs',r,8)) if v('Unit_Costs',r,8) is not None else '—'} |")
w("")
w("Human labour (the 2D confirm per plan plus output QA plus support) is **about 39–79% of base cost** for every deliverable except the 60s cinematic video, where Veo seconds dominate. "
  "LLM tokens cost cents per deliverable and do not matter.")
w("")
fr = find("Unit_Costs", "Retry / failure sensitivity") + 2
w("Retry/failure sensitivity (Base case). Cost per success = fixed + repeatable/(1 − f), and rejected attempts re-incur generation and QA:")
w("")
w("| Deliverable | £ @5% | £ @15% | £ @30% | £ @50% |")
w("|---|---|---|---|---|")
for r in range(fr, fr + 10):
    w(f"| {v('Unit_Costs',r,1)} | " + " | ".join(gbp(v('Unit_Costs', r, c), 2) for c in range(6, 10)) + " |")
w("")
w("## Pricing recommendation")
pp = {k: find("Pricing", k) for k in ["List price", "Included credits", "Effective price per", "Overage price", "Setup fee (£", "Per-development fee, extra", "Hosted interactive", "SSO connections"]}
base_blk = find("Pricing", "BASE cost case")
high_blk = find("Pricing", "HIGH cost case")
g100b = find("Pricing", "  Gross margin at 100%", base_blk)
gutb = find("Pricing", "  Gross margin at Base", base_blk)
ogb = find("Pricing", "  Overage gross", base_blk)
g100h = find("Pricing", "  Gross margin at 100%", high_blk)
w("Keep the hypothesised price anchors, but sell **credits** instead of 'unlimited'. Every geometry-true output needs a human 2D confirm per plan, so single-image orders are labour-heavy and need a per-plan credit charge.")
w("")
w("| | Studio | Agency | Developer Pro | Enterprise |")
w("|---|---|---|---|---|")
P = lambda r, c: v("Pricing", r, c)
w("| List price / month | " + " | ".join(gbp(P(pp["List price"], c)) for c in range(2, 6)) + " |")
w("| Included credits / month | " + " | ".join(num(P(pp["Included credits"], c), 0) for c in range(2, 6)) + " |")
w("| Effective £ per included credit | " + " | ".join(gbp(P(pp["Effective price per"], c), 2) for c in range(2, 6)) + " |")
w("| Overage £ / credit | " + " | ".join(gbp(P(pp["Overage price"], c), 2) for c in range(2, 6)) + " |")
w("| Setup fee | " + " | ".join(gbp(P(pp["Setup fee (£"], c)) for c in range(2, 6)) + " |")
w("| Extra development fee per plan type | " + " | ".join(gbp(P(pp["Per-development fee, extra"], c)) for c in range(2, 6)) + " |")
w("| Hosted interactive properties / SSO | " + " | ".join(f"{int(P(pp['Hosted interactive'], c))} / {int(P(pp['SSO connections'], c))}" for c in range(2, 6)) + " |")
w("| GM at 100% of included usage (Base cost) | " + " | ".join(pct(P(g100b, c)) for c in range(2, 6)) + " |")
w("| GM at 65% utilisation (Base cost) | " + " | ".join(pct(P(gutb, c)) for c in range(2, 6)) + " |")
w("| GM at 100% usage (High cost) | " + " | ".join(pct(P(g100h, c)) for c in range(2, 6)) + " |")
w("| Overage GM (Base cost) | " + " | ".join(pct(P(ogb, c)) for c in range(2, 6)) + " |")
w("")
hy = find("Pricing", "Hypothesis") + 1
w("Hypothesis price points tested against the same usage mixes (GM at 100% usage, Base / High cost): "
  + "; ".join(f"{P(r,1)} {pct(P(r,4))} / {pct(P(r,5))}" for r in range(hy, hy + 6)) + ".")
w("")
gr = find("Pricing", "Target gross margin")
w("Monthly price needed for a target gross margin at 100% of included usage:")
w("")
w("| Tier | 60% (Base) | 70% | 80% | 85% | 60% (High) | 70% | 80% | 85% |")
w("|---|---|---|---|---|---|---|---|---|")
tr = find("Pricing", "Subscription tiers at 100%", gr) + 1
for r in range(tr, tr + 4):
    w(f"| {P(r,1)} | " + " | ".join(gbp(P(r, c)) for c in range(2, 10)) + " |")
dr = find("Pricing", "One-off deliverables", gr) + 1
w("")
w("One-off or per-development prices needed for 70% / 80% GM at Base cost: "
  + "; ".join(f"{P(r,1)} {gbp(P(r,3),0)} / {gbp(P(r,4),0)}" for r in (dr + 2, dr + 3, dr + 7, dr + 8, dr + 9)) + ".")
devr = find("Pricing", "  Fee charged")
w(f"An extra 50-unit development billed at {gbp(P(devr,2))} (6 plan types × £350) earns a {pct(P(devr+2,3))} margin at Base cost and {pct(P(devr+3,3))} at High cost. "
  "At High cost, price per plan type should rise to £500–600, which is the top of the competitor anchor of £300–£600 per plot type.")
w("")
w("The recommended structure has five parts: (1) subscription with included credits; (2) overage priced above the included-credit rate; (3) setup fees that cover onboarding labour (Agency £500, Developer Pro £1,500, Enterprise £7,500); "
  "(4) per-development fees by plan type, because cost scales with plan types, not units; (5) WebGL viewers as the default, with pixel-streamed tours sold as a paid add-on, because streaming is the largest cost risk at High cost.")
w("")
w("## Scenario outputs")
w("| Metric | Bear M12 | Bear M24 | Bear M36 | Base M12 | Base M24 | Base M36 | Bull M12 | Bull M24 | Bull M36 |")
w("|---|---|---|---|---|---|---|---|---|---|")
fmtmap = {"Subscription MRR": gbp, "ARR (£)": gbp, "Paying customers": lambda x: num(x, 1), "Gross margin (month)": pct,
          "  memo: GM": pct, "Monthly EBITDA": gbp, "Cash balance": gbp, "Headcount": lambda x: num(x, 0)}
names = {"Subscription MRR": "Subscription MRR", "ARR (£)": "ARR", "Paying customers": "Customers", "Gross margin (month)": "Gross margin",
         "  memo: GM": "GM incl. support/CS salaries", "Monthly EBITDA": "Monthly EBITDA", "Cash balance": "Cash balance (cumulative)", "Headcount": "Headcount"}
for k in fmtmap:
    w(f"| {names[k]} | " + " | ".join(fmtmap[k](B(k, c)) for c in range(2, 11)) + " |")
w("")
w("| Scenario metric | Bear | Base | Bull |")
w("|---|---|---|---|")
for k, lab, f in [("First month with EBITDA", "First month EBITDA ≥ 0", mon), ("Sustained break-even", "Sustained break-even month", mon),
                  ("First month cash goes", "Cash first negative", mon), ("Lowest cash", "Lowest cash balance", gbp),
                  ("Month of lowest", "Month of lowest cash", mon), ("External funding", "External funding needed", gbp),
                  ("Cumulative revenue", "36-month cumulative revenue", gbp), ("Gross margin, 36", "36-month gross margin", pct),
                  ("Month first support", "Support/QA hire #1 starts", mon), ("Month Engineer", "Engineer #1 starts", mon),
                  ("Month AE", "AE #1 starts", mon), ("Month first Enterprise", "First enterprise signing (fractional)", mon)]:
    w(f"| {lab} | " + " | ".join(f(S(k, c)) for c in (2, 3, 4)) + " |")
w("")
w("## Milestones (first month reached; subscription MRR basis)")
mr = find("Milestones", "£10k MRR")
w("| Milestone | Bear | Base | Bull |")
w("|---|---|---|---|")
for r in range(mr, mr + 7):
    w(f"| {v('Milestones',r,1)} | " + " | ".join(mon(v("Milestones", r, c)) for c in (3, 4, 5)) + " |")
w("")
xr = find("Milestones", "£10k MRR", mr + 8)
w("Example customer mix needed for each milestone, using the Base month-36 MRR mix and ARPA (Studio / Agency / Developer Pro / Enterprise = total): "
  + "; ".join(f"{v('Milestones',r,1)}: {num(v('Milestones',r,3),0)} / {num(v('Milestones',r,4),0)} / {num(v('Milestones',r,5),0)} / {num(v('Milestones',r,6),0)} = {num(v('Milestones',r,7),0)}"
              for r in range(xr, xr + 7)) + ". "
  "For scale, the UK has about 500 medium developers, 200 off-plan sellers and 25 nationals, and the researcher's SOM is about £1.25m ARR by year 3. £5m ARR would need about 23 enterprise accounts and about 92 Developer Pro accounts at today's mix. That is roughly the whole national-housebuilder universe plus about 18% of medium developers, so it is not credible as UK-only within 36 months.")
w("")
w("## Solo-founder capacity")
cr = find("Summary", "Founder hours available")
w(f"The founder is assumed to work about 180 h/month, with 35% ({num(v('Summary',cr,2),0)} h) available for delivery operations: 2D confirm, QA, support and onboarding. The rest goes to building the product (6–15 engineering-months per technical_feasibility), selling and admin. "
  "Capacity = available minutes ÷ human minutes per unit, at Base cost and including rejects:")
w("")
w("| Unit | Human minutes | Max per month (founder only) |")
w("|---|---|---|")
for r in range(cr + 2, cr + 12):
    w(f"| {v('Summary',r,1)} | {num(v('Summary',r,2),1)} | {num(v('Summary',r,3),1)} |")
w("")
w(f"So a solo founder can deliver about {num(v('Summary',cr+3,3),0)} 2-bed-flat packages or about {num(v('Summary',cr+4,3),0)} 4-bed houses a month. That is roughly {num(v('Summary',cr+8,3),0)} Studio, {num(v('Summary',cr+9,3),0)} Agency, {num(v('Summary',cr+10,3),0)} Developer Pro or {num(v('Summary',cr+11,3),1)} Enterprise accounts on their own. "
  f"In the Base scenario, delivery hours exceed founder capacity in month {S('First month human',3)}, at about {num(S('Customers in that',3),0)} customers. The first support/QA hire, triggered at 20 customers, starts in month {S('Month first support',3)}. "
  "QA overflow is costed in COGS at contractor rates throughout.")
w("")
w("## Caveats")
w("- Customers are fractional expected values. The model ignores tax, VAT timing, R&D tax credits, working capital and capex. MRR excludes overage, setup and per-development fees.")
w("- The Base ~70% gross margin is above ICONIQ's ~45–55% average for AI products because COGS here excludes support/CS salaries (the memo row adds them back) and assumes offshore QA. If customers expect studio-style revisions, the business behaves like a tech-enabled service (40–55%).")
w("- Prices, allowances, QA minutes, rejection rates, conversion and churn are hypotheses. Replace them with pilot data (time trials on 50 plans; Van Westendorp pricing tests; 3 paid developer pilots).")
open(OUT, "w").write("\n".join(L) + "\n")
print("wrote", OUT)
